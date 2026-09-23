"""Create a presentation copy from the existing deck and verified GIF catalogue."""
from pathlib import Path
import json, hashlib, io, zipfile, re, copy
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import build_slide_deck as b
import build_stage_report as report

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / '阶段仿真实现汇报.pptx'
OUT = ROOT / '阶段仿真实现汇报_GIF优化版.pptx'
ROWS = json.loads((ROOT/'simulation_gifs_20260923/case_manifest.json').read_text(encoding='utf-8'))
INDEX = {(r['group'], r['name']): r for r in ROWS}
USED = []
prs = Presentation(SOURCE)
source_hash = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
while len(prs.slides) > 6:
    entry = prs.slides._sldIdLst[-1]
    prs.part.drop_rel(entry.rId)
    prs.slides._sldIdLst.remove(entry)

def text(slide, content, x,y,w,h,size=17,bold=False,color=None):
    f=b.textbox(slide,Inches(x),Inches(y),Inches(w),Inches(h))
    f.word_wrap=True
    b.paragraph(f,content,size,bold,color or b.INK,first=True,space_after=0)
    return f

def title(slide, heading, sub=''):
    b.accent(slide,top=Inches(.38),height=Inches(.52))
    text(slide,heading,.85,.30,11.8,.64,27,True)
    if sub:text(slide,sub,.85,1.03,11.8,.55,14,color=b.MUTED)

def notes(slide, content): b.set_notes(slide,content)

def gif_page(stage, heading, group, name, bullets, caution):
    r=INDEX[group,name]
    s=b.blank(prs); title(s,heading,stage+'  |  已保存结果回放；物理时间见动画，不等于播放时长')
    y=1.93
    for head,body in bullets:
        text(s,head,.85,y,2.78,.38,18,True,b.BRAND)
        text(s,body,.85,y+.42,2.78,1.02,16)
        y+=1.48
    path=r['gif']; bw,bh=Inches(8.75),Inches(5.35)
    w,h=b.fit(path,bw,bh)
    s.shapes.add_picture(path,Inches(3.90)+(bw-w)//2,Inches(1.60)+(bh-h)//2,width=w,height=h)
    text(s,caution,.85,7.04,11.7,.31,10.5,color=b.MUTED)
    notes(s,'\n'.join([heading,*(a+'：'+v for a,v in bullets),caution,
         '算例：'+name,'来源：'+'；'.join(r.get('sources',[])),
         '原始 GIF：'+path,'限制：'+'；'.join(r.get('caveats',[])),
         '合并动画的各段可能重新从 t=0 开始，不表示连续加载历史。不同动画色标可能不同。',
         'GIF 原样嵌入；原始分段最多51个保存帧，未补造求解结果。']))
    USED.append({'slide':len(prs.slides),'title':heading,**r})

def evidence(heading, lead, keys, extra=''):
    s=b.blank(prs);title(s,heading,lead)
    n=len(keys)
    boxes=([(0.85,1.80,11.7,4.95)] if n==1 else
           [(0.75,2.05,5.9,4.70),(6.8,2.05,5.8,4.70)] if n==2 else
           [(0.75,1.75,6.9,4.98),(7.85,1.7,4.8,2.42),(7.85,4.26,4.8,2.42)])
    for key,(x,y,w,h) in zip(keys,boxes):
        p=ROOT/report.FIGURES[key];iw,ih=b.fit(str(p),Inches(w),Inches(h-.40))
        s.shapes.add_picture(str(p),Inches(x)+(Inches(w)-iw)//2,Inches(y)+(Inches(h-.4)-ih)//2,width=iw,height=ih)
        text(s,b.caption_for(key),x,y+h-.34,w,.4,11,color=b.MUTED)
    if extra:text(s,extra,.85,7.04,11.7,.32,10.5,color=b.MUTED)
    notes(s,lead+'\n'+extra+'\n图源：\n'+'\n'.join(str(ROOT/report.FIGURES[k]) for k in keys))

# Keep the original theme and foundation slides, edit only the cover and overview text.
cover=prs.slides[0]
for sh in list(cover.shapes):
    el=sh._element;el.getparent().remove(el)
title(cover,'落球冲击、层间损伤与主动导波检测')
text(cover,'仿真实现阶段汇报',.85,2.05,11.5,.85,38,True)
text(cover,'GIF 优化版 · 代表性算例与数值证据',.85,3.03,11.5,.65,23,color=b.BRAND)
text(cover,'文献实验基础与八阶段研究路线保留；主要动画单页放大。\n正文说明观察重点，备注记录来源、工况与适用范围。',.85,4.20,11.4,1.3,19)
text(cover,'2026-09-23  ·  动画为已有计算回放，尚未完成实体实验验证',.85,6.57,11.5,.5,14,color=b.MUTED)
notes(cover,'本版本由旧PPT副本改编，保留研究基础和八阶段主线。新增13个代表性GIF，原始文件未改动。其余动画在原素材目录。GIF须在支持该格式的PowerPoint放映中查看；静态导出仅显示一帧。')
for slide in list(prs.slides)[1:6]:
    for shape in slide.shapes:
        if shape.has_text_frame:
            if 'Fig.' in shape.text and shape.height < Inches(.4):
                shape.height=Inches(.43)
            for p in shape.text_frame.paragraphs:
                for run in p.runs:
                    run.text=run.text.replace('下一页起，按顺序讲这八步各做了什么、为什么这么做','后续按八阶段展示代表性动画与计算依据').replace('从下一页开始','后续技术部分展开')

gif_page('阶段一｜薄板','薄板落球：整体弯曲随时间变化','61_thin_plate','order46',[
 ('模型','500×400×2 mm 板；8 mm 球；落高160 mm。'),
 ('看什么','先看顶面位移，再看响应时程。w 向下为正。'),
 ('可信范围','每方向46个基函数；整体位移较稳，局部接触仍敏感。')],
 '多段回放会重新计时；薄板模型不用于真实层间损伤预测。')
evidence('薄板模型的空间收敛','38→46阶：位移变化约0.6%，接触力仍变化约5.7%。',['s1_plate_modal'],
 '46²=2116个空间自由度；方板参考35.99为无量纲基频。')
gif_page('阶段二｜三维冲击','局部接触：刚球与板的响应','62_custom_3d','submodel',[
 ('两级求解','局部细网格求接触力，再驱动整板模型。'),('观察重点','接触区、总接触力与节点响应的演变。'),('结果范围','约32 N接触峰值；与薄板差异不等于实测误差。')],
 '此局部模型无损伤演化；局部场与后续全局响应不能混作同一网格解。')
gif_page('阶段二｜远场响应','传感器读数：应变随位置和方向变化','62_custom_3d','farfield_200x160x4',[
 ('网格','整板200×160×4，用同一接触力时程加载。'),('观察重点','不同通道的到达时刻、幅值及应变方向。'),('物理量','显示机械应变和保存测点位移，尚非PZT电压。')],
 '只有已保存测点时程的动画，不能重建未保存的整板全场。')
evidence('三维模型：接触差异、网格与方向','局部模型约311 μm，收敛全局模型约313 μm；方向投影可相差1.2–4.3倍。',
 ['s2_plate_vs_solid','s2_farfield_convergence','s2_sensor_direction'],'两模型接触表示不同；尚未完全分离差异来源，也没有实体对照。')
gif_page('阶段三｜传播验证','二维导波：已保存波场与传播时程','60_custom_wave','4000_32_healthy',[
 ('模型','4000×32网格，板厚1.72 mm；无损基线。'),('看什么','波形沿中心线移动，时空图呈现传播轨迹。'),('验证目标','在指定频段、方向上提取主导A0相速度。')],
 '动画中的波包移动不能直接代替相速度提取；S0与群速度未完成同等验证。')
evidence('A0 相速度与解析频散对照','100 kHz：三档网格结果趋稳，与1286.4 m/s解析值偏差约0.48%–0.68%。',
 ['s3_dispersion_analytic','s3_dispersion_check'],'比较需固定材料、厚度、频率、模态及传播方向。')
gif_page('阶段四｜Abaqus核验','Abaqus 线源：三维实体中的导波传播','30_abaqus_validation','ugw_line',[
 ('核验方法','统一输入设置，比较独立求解器的波速及波形。'),('看什么','中心线传播、时空图和模型动能历史。'),('结论边界','波速接近解析；完整波形仍有网格敏感性。')],
 '线源验证与三维点源不同；GIF说明过程，具体误差见下一页。')
evidence('求解器对照与单元选择','自研1285.6 m/s（−0.06%）；Abaqus C3D8 1288.5 m/s（+0.16%）。',
 ['s4_cross_solver','s4_abaqus_speed','s4_reduced_integration'],'0.451>0.322、0.365<0.862：不能将两测点差异都归因于网格；C3D8R结论限于已测设置。')

s=b.blank(prs);title(s,'阶段五：材料、单元与界面路线','最终采用C3D8实体层＋层间表面黏聚；层内判据仅诊断。')
items=[('C3D8＋原生Hashin','预处理拒绝','原生实现的适用单元限制'),('SC8R＋Hashin','能力试验通过','未进入最终主线'),('自定义VUMAT','本次未采用','测试时未配置编译环境'),('实体层内判据','仅诊断','当前算例不退化层内刚度'),('表面黏聚','最终采用','层间损伤引起界面刚度退化')]
t=s.shapes.add_table(6,3,Inches(.85),Inches(1.95),Inches(11.7),Inches(3.75)).table
for row,vals in enumerate([('路线','结果','含义')]+items):
 for col,v in enumerate(vals):
  c=t.cell(row,col);c.text=v;c.fill.solid();c.fill.fore_color.rgb=b.BRAND if row==0 else b.WHITE
  for p in c.text_frame.paragraphs:
   for run in p.runs:b.style_run(run,16,row==0,b.WHITE if row==0 else b.INK)
t.columns[0].width=Inches(4);t.columns[1].width=Inches(2.5);t.columns[2].width=Inches(5.2)
text(s,'下一组动画展示关键词能力探针。能运行与完成真实材料标定是不同证据。',.85,6.30,11.7,.65,17)
notes(s,b.SLIDES[4]['notes'])
gif_page('阶段五｜能力探针','连续壳 Hashin：可运行路线的单元试验','40_material_probes','axial_sc8r_hashin',[
 ('试验目的','核对SC8R与Hashin组合能否运行。'),('观察重点','加载过程中保存的节点运动与响应。'),('范围','属于单元/关键词探针，最终大模型未采用该路线。')],
 '位移随加载变化本身不能证明损伤演化或材料参数已经标定。')
gif_page('阶段五｜界面探针','表面黏聚：叠层界面的分离过程','40_material_probes','surfcoh_stack',[
 ('试验目的','检查表面黏聚接触的界面响应。'),('观察重点','两侧节点相对位移与已保存时间历程。'),('模型含义','界面可承载并发生损伤；与层内破坏不同。')],
 '小算例仅用于能力检查，不代表实际试件或实体实验。')

for name,en,dmd,rad in [('imp_lo','0.100','0.0966','0.17–0.25'),('imp_mid','0.300','0.9267','0.53–0.78'),('imp_hi','0.794','15.11','2.13–3.13')]:
 gif_page('阶段六｜单次冲击',en+' J 冲击：接触、变形与损伤耗能','01_current_impact',name,[
  ('试件','100×100×2 mm；8层同向材料轴；中心网格0.15 mm。'),
  ('损伤指标','ALLDMD='+dmd+' mJ；能量等价圆半径'+rad+' mm。'),
  ('看什么','球与板截面、顶面U3和总损伤耗能。不同GIF色标可不同。')],
  'K=2.37×10¹³ N/m³；等价圆不是实际分层边界，局部接触应力尚未收敛。')
evidence('三档冲击标定与局部网格','0.300 J：接触峰值约569 N、挠度约782 μm、损伤耗能0.9267 mJ。',
 ['s6_impact_energies','s6_graded_mesh','s6_impact_history'],'等价面积由多界面总耗能换算；参数为文献替代值，尚无实物标定。')
for name,label,metric in [('wav_base','无损基线','镜像V3归一化差：0.0012–0.0049。'),('wav_d03_r08','0.8 mm 理想脱黏','R1/R2波形差：0.0547/0.0463。'),('wav_d08_r31','3.1 mm 理想脱黏','R1/R2波形差：1.50/1.33。')]:
 gif_page('阶段七｜主动检测',label+'：导波传播与接收响应','02_current_wave',name,[
  ('激励','100 kHz、5周期；压力片半径3 mm；1激励＋8接收。'),
  ('定量对照',metric+'比较窗口0–90 μs。'),
  ('观察重点','波前、中心线响应及反射；晚时刻不能混作直达波。')],
  '静止初态；理想脱黏未逐点继承冲击损伤。0.8 mm档仍有4/5界面映射差异。')
evidence('导波检测：波形差与数值基线','小圆斑引起约5%的归一化波形差；大圆斑引起明显形状和相位改变。',
 ['s7_wave_waveforms','s7_sensor_layout','s7_wave_floor'],'DI不是识别准确率；负相关不等于整段严格反相；数值镜像基线不是实验电子噪声。')
evidence('阶段八：界面稳定性与参数敏感性','显式给出黏聚刚度后，已测试算例的镜像一致性改善；失稳机理仍未完全确定。',
 ['s8_instability'],'初始步长相同不证明接触刚度完全不参与稳定性估计；三档等价面积下降约73%、51%、50%。')
s=b.blank(prs);title(s,'成果范围与下一步验证')
for y,h,body in [(1.7,'已形成的数值证据','整体响应、A0传播、求解器对照，以及理想脱黏的导波敏感性。'),(3.1,'尚未闭环的部分','真实损伤逐点传递、局部接触应力收敛、PZT机电响应及独立实体实验验证。'),(4.5,'与厦大实验的关系','继承冲击监测与健康监测方法；试件、冲击源、传感链路和验证任务尚未一一对应。')]:
 text(s,h,.85,y,11.5,.45,22,True,b.BRAND);text(s,body,.85,y+.56,11.5,.80,20)
notes(s,'未将正式累积序列的第1次冲击冒充三次累积完成；本版本仅选取当前单次标定与阶段核验动画。'+ '\n'.join(a+'：'+v for a,v in b.CAVEATS))
s=b.blank(prs);title(s,'动画使用与结果追溯')
for j,(h,body) in enumerate([('原始GIF已嵌入','主汇报选取13个代表性算例，全部动画仍在 simulation_gifs_20260923。'),('播放与导出','使用PowerPoint放映查看动画。静态PDF或缩略图仅显示一帧。'),('来源与限制','每个动画页备注记录原文件、数据来源、工况与限制；合并段落可重置物理时间。'),('复算资源','旧PPT保留。原论文、ODB及GIF为本地外部资源；新版本没有重新运行求解。')]):
 y=1.65+j*1.22;text(s,h,.85,y,11.5,.4,21,True,b.BRAND);text(s,body,.85,y+.45,11.5,.65,17)
notes(s,'生成脚本：review/build_review_gif_optimized.py\n原PPT：'+str(SOURCE)+'\n来源清单：review/simulation_gifs_20260923/case_manifest.json\n逐页动画映射：review/GIF优化版_素材映射.json')

# Every technical page, including evidence pages, carries its stage in the main title.
stage_ranges=[(7,8,'一'),(9,11,'二'),(12,13,'三'),(14,15,'四'),
              (16,18,'五'),(19,22,'六'),(23,26,'七'),(27,27,'八')]
for first,last,number in stage_ranges:
    for slide_number in range(first,last+1):
        slide=prs.slides[slide_number-1]
        heading=next(sh for sh in slide.shapes if sh.has_text_frame and sh.top<Inches(.5) and sh.left>Inches(.7))
        content=re.sub(r'^阶段[一二三四五六七八][：:｜|]\s*','',heading.text)
        heading.text_frame.clear()
        p=heading.text_frame.paragraphs[0]
        r=p.add_run();r.text='阶段'+number+'｜';b.style_run(r,27,True,RGBColor.from_string('B91C1C'))
        r=p.add_run();r.text=content;b.style_run(r,25,True,b.INK)
important=re.compile(r'尚未完成实体实验验证|尚无实物标定|没有实体对照|不代表实际试件|未逐点继承冲击损伤|理想脱黏|不能重建未保存的整板全场|尚非PZT电压|不退化层内刚度|最终采用|仅诊断|未采用|局部接触应力尚未收敛|失稳机理仍未完全确定|尚未闭环|独立实体实验验证|C3D8实体层＋层间表面黏聚|ALLDMD|DI不是识别准确率|0\.100 J|0\.300 J|0\.794 J|0\.8 mm|3\.1 mm|0\.0966|0\.9267|15\.11|0\.0547/0\.0463|1\.50/1\.33|1286\.4|1285\.6|1288\.5|0\.6%|5\.7%|73%、51%、50%|569 N|782 μm|13个代表性算例|100 kHz|5周期')
highlighted=0
for s in prs.slides:
    frames=[]
    for sh in s.shapes:
        if sh.has_text_frame:frames.append(sh.text_frame)
        if sh.has_table:frames.extend(c.text_frame for row in sh.table.rows for c in row.cells)
    for frame in frames:
        for p in frame.paragraphs:
            for run in list(p.runs):
                size=run.font.size.pt if run.font.size else 16
                if size<10:continue
                matches=list(important.finditer(run.text))
                if not matches:continue
                original=run._r; parent=original.getparent();idx=parent.index(original);pos=0
                segments=[]
                for m in matches:
                    if m.start()>pos:segments.append((run.text[pos:m.start()],False))
                    segments.append((m.group(),True));pos=m.end()
                if pos<len(run.text):segments.append((run.text[pos:],False))
                for value,red in segments:
                    new=copy.deepcopy(original);new.find('{http://schemas.openxmlformats.org/drawingml/2006/main}t').text=value
                    parent.insert(idx,new);idx+=1
                    if red:
                        from pptx.text.text import _Run
                        rr=_Run(new,p);rr.font.bold=True;rr.font.color.rgb=RGBColor.from_string('B91C1C');highlighted+=1
                parent.remove(original)
for i,s in enumerate(prs.slides,1):
    frame=text(s,f'{i:02d} / {len(prs.slides):02d}',12.12,7.31,.85,.17,8,color=b.MUTED)
    frame.margin_left=frame.margin_right=frame.margin_top=frame.margin_bottom=0
    frame.word_wrap=False
prs.save(OUT)
problems=[]
for i,s in enumerate(prs.slides,1):
 for sh in s.shapes:
  if sh.left<0 or sh.top<0 or sh.left+sh.width>prs.slide_width or sh.top+sh.height>prs.slide_height:problems.append((i,sh.name))
with zipfile.ZipFile(OUT) as z:
 gifs=[n for n in z.namelist() if n.startswith('ppt/media/') and n.endswith('.gif')]
 gif_checks=[{'part':n,'frames':Image.open(io.BytesIO(z.read(n))).n_frames} for n in gifs]
assert not problems, problems
assert len(gif_checks)==len(USED)==13
assert all(r['frames']>1 for r in gif_checks)
assert hashlib.sha256(SOURCE.read_bytes()).hexdigest()==source_hash
(ROOT/'GIF优化版_素材映射.json').write_text(json.dumps({'source':str(SOURCE),'source_sha256':source_hash,'output':str(OUT),'slides':len(prs.slides),'selected':USED,'embedded_gifs':gif_checks},ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'output':str(OUT),'slides':len(prs.slides),'GIFs':len(gifs),'red_bold_spans':highlighted,'MB':OUT.stat().st_size/1e6,'overflow':problems},ensure_ascii=False))
