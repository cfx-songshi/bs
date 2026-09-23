"""Build the 28-slide explanation report without changing its PPT."""
from pathlib import Path
import re, json, hashlib
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pptx import Presentation

ROOT=Path(__file__).resolve().parent
SOURCE=ROOT/'阶段仿真实现汇报_GIF优化版_逐页讲解报告.md'
OUT=SOURCE.with_suffix('.docx')
PPT=ROOT/'阶段仿真实现汇报.pptx'
content=SOURCE.read_text(encoding='utf-8')
pages=[int(x) for x in re.findall(r'^## 第(\d+)页',content,re.M)]
assert pages==list(range(1,29)),pages
assert len(Presentation(PPT).slides)==28
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.75)
sec.left_margin=sec.right_margin=Inches(.85)
sec.header_distance=sec.footer_distance=Inches(.3)

def font(run,size=11,bold=False,color='000000'):
    run.font.name='Microsoft YaHei';run.font.size=Pt(size);run.font.bold=bold
    run.font.color.rgb=RGBColor.from_string(color)
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')

for name,size in [('Normal',11),('Title',22),('Heading 1',15),('Heading 2',12)]:
    st=doc.styles[name];st.font.name='Microsoft YaHei';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
    st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')
    st.paragraph_format.line_spacing=1.35
    st.paragraph_format.space_after=Pt(6)
    if name!='Normal':st.paragraph_format.keep_with_next=True
doc.styles['Heading 1'].paragraph_format.space_before=Pt(15)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(8)
doc.styles['Normal'].paragraph_format.widow_control=True

p=sec.header.paragraphs[0];font(p.add_run('阶段仿真实现汇报  GIF优化版逐页讲解'),9)
p=sec.footer.paragraphs[0];p.alignment=2
font(p.add_run('讲解报告  '),9)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');p._p.append(field)

for line in content.splitlines():
    line=line.strip()
    if not line:continue
    if line.startswith('# '):
        p=doc.add_paragraph(style='Title');font(p.add_run(line[2:]),22,True);continue
    if line.startswith('## '):
        p=doc.add_paragraph(style='Heading 1');font(p.add_run(line[3:]),15,True);continue
    if re.fullmatch(r'\*\*[^*]+\*\*',line):
        p=doc.add_paragraph(style='Heading 2');font(p.add_run(line[2:-2]),12,True);continue
    p=doc.add_paragraph()
    if line.startswith('- '):
        line='• '+line[2:];p.paragraph_format.left_indent=Inches(.15)
        p.paragraph_format.first_line_indent=Inches(-.12)
    for part in re.split(r'(\*\*.*?\*\*)',line):
        if part.startswith('**') and part.endswith('**'):font(p.add_run(part[2:-2]),11,True,'B91C1C')
        else:font(p.add_run(part),11)
    p.paragraph_format.line_spacing=1.35
    p.paragraph_format.space_after=Pt(6)

doc.core_properties.title='阶段仿真实现汇报逐页讲解报告'
doc.core_properties.subject='对应28页GIF优化版 含八阶段讲解 读图说明与结论边界'
# Remove the default template's decorative paragraph borders.
for elem in list(doc.styles.element.iter(qn('w:pBdr'))) + list(doc.element.iter(qn('w:pBdr'))):
    elem.getparent().remove(elem)
doc.save(OUT)
(ROOT/'gif_explanation_version.json').write_text(json.dumps({'ppt':str(PPT),'ppt_sha256':hashlib.sha256(PPT.read_bytes()).hexdigest(),'report':str(OUT),'slide_sections':pages,'markdown_chars':len(content)},ensure_ascii=False,indent=2),encoding='utf-8')
print('Wrote DOCX; 28 slide sections verified; chars:',len(content))
